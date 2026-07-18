"""Renderiza una vista previa PNG de las hojas de un XLSX para QA visual local."""

from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw, ImageFont


def _fuente(tam: int, negrita: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    nombres = (
        ("C:/Windows/Fonts/arialbd.ttf", "DejaVuSans-Bold.ttf")
        if negrita
        else ("C:/Windows/Fonts/arial.ttf", "DejaVuSans.ttf")
    )
    for nombre in nombres:
        try:
            return ImageFont.truetype(nombre, tam)
        except OSError:
            continue
    return ImageFont.load_default()


def _rgb(valor: object, defecto: tuple[int, int, int]) -> tuple[int, int, int]:
    texto = str(valor or "")
    if len(texto) == 8:
        texto = texto[2:]
    if len(texto) == 6:
        try:
            return tuple(int(texto[i : i + 2], 16) for i in (0, 2, 4))
        except ValueError:
            pass
    return defecto


def _recortar(draw: ImageDraw.ImageDraw, texto: str, ancho: int, fuente: ImageFont.ImageFont) -> str:
    texto = texto.replace("\n", " ")
    if draw.textbbox((0, 0), texto, font=fuente)[2] <= ancho:
        return texto
    while texto and draw.textbbox((0, 0), texto + "…", font=fuente)[2] > ancho:
        texto = texto[:-1]
    return texto + "…"


def renderizar(path: Path, destino: Path, max_filas: int = 18) -> Path:
    wb = load_workbook(path, read_only=False, data_only=False)
    vistas: list[Image.Image] = []
    titulo_font = _fuente(22, True)
    for ws in wb.worksheets:
        n_filas = min(ws.max_row, max_filas)
        n_cols = ws.max_column
        anchos = []
        for col in range(1, n_cols + 1):
            ancho_excel = ws.column_dimensions[get_column_letter(col)].width or 12
            anchos.append(max(80, min(240, int(ancho_excel * 7.2))))
        alto_fila = 30
        margen, alto_titulo = 16, 46
        imagen = Image.new(
            "RGB", (sum(anchos) + 2 * margen, alto_titulo + n_filas * alto_fila + margen), "white"
        )
        draw = ImageDraw.Draw(imagen)
        draw.rectangle((0, 0, imagen.width, alto_titulo), fill=(30, 58, 95))
        draw.text((margen, 10), f"{path.name} · hoja {ws.title}", font=titulo_font, fill="white")
        y = alto_titulo
        for fila in range(1, n_filas + 1):
            x = margen
            for col, ancho in enumerate(anchos, start=1):
                celda = ws.cell(fila, col)
                relleno = _rgb(celda.fill.fgColor.rgb, (255, 255, 255))
                if celda.fill.fill_type is None:
                    relleno = (255, 255, 255)
                color = _rgb(celda.font.color.rgb if celda.font.color and celda.font.color.type == "rgb" else None, (25, 25, 25))
                fuente = _fuente(13, bool(celda.font.bold))
                draw.rectangle((x, y, x + ancho, y + alto_fila), fill=relleno, outline=(185, 190, 198))
                texto = "" if celda.value is None else str(celda.value)
                texto = _recortar(draw, texto, ancho - 10, fuente)
                draw.text((x + 5, y + 7), texto, font=fuente, fill=color)
                x += ancho
            y += alto_fila
        vistas.append(imagen)

    ancho = max(v.width for v in vistas)
    alto = sum(v.height for v in vistas) + 12 * (len(vistas) - 1)
    combinada = Image.new("RGB", (ancho, alto), (225, 229, 235))
    y = 0
    for vista in vistas:
        combinada.paste(vista, (0, y))
        y += vista.height + 12
    destino.parent.mkdir(parents=True, exist_ok=True)
    combinada.save(destino)
    return destino


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("archivos", nargs="+", type=Path)
    parser.add_argument("--destino", type=Path, default=Path("pyme_hibrida/reconstruccion_smote/qa_excel"))
    args = parser.parse_args()
    for archivo in args.archivos:
        salida = renderizar(archivo, args.destino / f"{archivo.stem}.png")
        print(salida)


if __name__ == "__main__":
    main()
