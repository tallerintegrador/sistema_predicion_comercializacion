"""Genera el Documento de Pruebas de Caja Negra del sistema SPC.

Replica el formato del entregable de referencia (`final/Pruebas de caja negra.docx`):
portada + tabla índice + una tabla 2-columnas por escenario, con la misma paleta de
colores y estructura de secciones. Los datos de los 16 escenarios son la única fuente
de verdad y desde ellos se emiten tanto el `.docx` como el `.md` fuente.

Uso:
    python scripts/gen_pruebas_caja_negra.py

Salida:
    final/Pruebas de Caja Negra SPC.docx
    final/Pruebas de Caja Negra SPC.md
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor

# --- Paleta (idéntica al documento de referencia) -------------------------------
AZUL_OSCURO = "1F4E79"   # cabecera de escenario y de la tabla índice
AZUL_MEDIO = "2E75B6"    # cabeceras de sección (DATOS DE ENTRADA, ...)
AZUL_CLARO = "D6E4F0"    # celdas de etiqueta (columna izquierda)
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)
GRIS_TXT = RGBColor(0x22, 0x22, 0x22)

HARDWARE_SW = (
    "PC con navegador Google Chrome (última versión). Backend FastAPI (Uvicorn) "
    "ejecutado localmente o en Render. Frontend React 19 + Vite. Base de datos "
    "PostgreSQL (Supabase / Neon)."
)
RESPONSABLE = "Equipo SPC"
ESTADO = "Pendiente"


# ================================================================================
# Datos de los 16 escenarios (única fuente de verdad)
# ================================================================================
ESCENARIOS: list[dict] = [
    {
        "num": 1,
        "id_hu": "HU-1.1: Autenticación",
        "nombre": "Inicio y Cierre de Sesión",
        "modulo": "Autenticación",
        "tecnica": "Partición de equivalencias y análisis de valores borde",
        "descripcion": (
            "Verificar que el sistema permita iniciar y cerrar sesión con credenciales "
            "válidas e inválidas, emita un token de sesión firmado y proteja el acceso a "
            "las secciones internas."
        ),
        "entorno": (
            "Pantalla de inicio de sesión de SPC con los campos usuario (id) y contraseña, "
            "y el botón Iniciar Sesión. El cierre de sesión se realiza desde el menú de la "
            "cabecera."
        ),
        "parametros": ["Campo: Usuario (id)", "Campo: Contraseña", "Botón: Iniciar Sesión", "Botón: Cerrar Sesión"],
        "respuesta_modulos": (
            "Se invoca POST /auth/login, que valida las credenciales contra el hash "
            "almacenado y emite un token; GET /auth/me devuelve la identidad y permisos."
        ),
        "condiciones": [
            "Se ingresa usuario y contraseña correctos de una cuenta activa.",
            "Se ingresa usuario correcto y contraseña incorrecta.",
            "Se ingresan las credenciales de una cuenta desactivada (is_active = false).",
            "Se dejan vacíos los campos de usuario y contraseña y se presiona Iniciar Sesión.",
            "Se intenta acceder a una sección interna con el token expirado o ausente.",
            "El usuario con sesión activa presiona Cerrar Sesión.",
        ],
        "resultados": [
            "El sistema redirige al panel principal (o al onboarding si es el primer ingreso). La sesión queda activa.",
            "El sistema muestra un mensaje genérico «Id o contraseña incorrectos» y no permite el acceso.",
            "El sistema rechaza el acceso con el mismo mensaje genérico, sin revelar que la cuenta existe.",
            "El sistema muestra validaciones indicando que los campos son obligatorios; no llama a la API.",
            "El sistema responde 401 y redirige al login sin exponer la funcionalidad protegida.",
            "El sistema invalida el token, limpia la sesión y regresa al login.",
        ],
        "metodo": (
            "Partición de equivalencias: credenciales válidas vs inválidas. Análisis de "
            "valores borde: campos vacíos, cuenta inactiva y token expirado."
        ),
        "modulos_sistema": "auth.py (login, me), service/seguridad.py (crear_token / verificar_token), LoginPage.tsx, AuthContext.tsx.",
        "procedimiento": [
            "Navegar a la URL del sistema SPC.",
            "Ejecutar cada condición ingresando los datos descritos.",
            "Presionar Iniciar Sesión y observar la respuesta.",
            "Para la condición 6: estando autenticado, presionar Cerrar Sesión.",
            "Registrar el resultado y capturar la pantalla.",
        ],
        "dependencias": ["Debe existir al menos una cuenta registrada.", "El backend debe estar en ejecución."],
    },
    {
        "num": 2,
        "id_hu": "HU-1.2: Restablecer contraseña",
        "nombre": "Restablecer Contraseña por Correo",
        "modulo": "Autenticación",
        "tecnica": "Tabla de decisión",
        "descripcion": (
            "Verificar el flujo de recuperación de contraseña por correo: solicitud del "
            "enlace, validez del token de un solo uso y fijación de la nueva contraseña."
        ),
        "entorno": (
            "Enlace «¿Olvidó su contraseña?» en el login y página de restablecimiento "
            "accesible por la URL /reset?token=… sin sesión previa."
        ),
        "parametros": ["Campo: Correo electrónico", "Enlace de restablecimiento (token)", "Campo: Nueva contraseña"],
        "respuesta_modulos": (
            "POST /auth/forgot genera un token ligado al hash actual y lo envía por correo; "
            "POST /auth/reset valida el token y actualiza la contraseña."
        ),
        "condiciones": [
            "Se solicita el enlace con un correo asociado a una cuenta activa.",
            "Se solicita el enlace con un correo que no existe o de una cuenta inactiva.",
            "Se abre el enlace y se fija una contraseña nueva válida.",
            "Se reutiliza el mismo enlace una segunda vez tras cambiar la contraseña.",
            "Se abre un enlace expirado (fuera del TTL de restablecimiento).",
        ],
        "resultados": [
            "El sistema envía el correo con el enlace y muestra una respuesta genérica de confirmación.",
            "El sistema muestra la misma respuesta genérica (no revela si el correo existe) y no envía correo.",
            "El sistema fija la nueva contraseña e invalida el token; permite iniciar sesión con ella.",
            "El sistema responde 400 «El enlace no es válido o ya expiró»; la contraseña no cambia.",
            "El sistema responde 400 e invita a solicitar un enlace nuevo.",
        ],
        "metodo": (
            "Tabla de decisión sobre las combinaciones (correo válido/ inválido) × (token "
            "vigente / usado / expirado) y el resultado esperado en cada rama."
        ),
        "modulos_sistema": "auth.py (forgot_password, reset_password), seguridad.py (crear_token_reset / verificar_token_reset), ResetPasswordPage.tsx.",
        "procedimiento": [
            "Desde el login, abrir «¿Olvidó su contraseña?» e ingresar el correo.",
            "Revisar el correo (o el log del backend si no hay SMTP) y abrir el enlace.",
            "Fijar la nueva contraseña y confirmar.",
            "Reintentar con el mismo enlace y con uno expirado.",
            "Registrar cada resultado y capturar la pantalla.",
        ],
        "dependencias": ["Debe existir una cuenta con correo registrado.", "SMTP configurado (o revisar el log del backend)."],
    },
    {
        "num": 3,
        "id_hu": "HU-1.3: Onboarding del negocio",
        "nombre": "Onboarding / Perfil de Negocio",
        "modulo": "Perfil",
        "tecnica": "Partición de equivalencias",
        "descripcion": (
            "Verificar que un usuario no administrador complete obligatoriamente el perfil "
            "de negocio en su primer ingreso y que el sistema valide cada opción."
        ),
        "entorno": (
            "Formulario de onboarding con nombre del negocio, sector, tamaño, región y "
            "moneda; se muestra tras el primer login de una cuenta no administradora."
        ),
        "parametros": ["Campo: Nombre del negocio", "Selección: Sector", "Selección: Tamaño", "Selección: Región", "Selección: Moneda"],
        "respuesta_modulos": (
            "GET /profile/options puebla los selectores; PUT /profile valida cada valor "
            "contra el conjunto permitido, persiste el perfil y marca onboarding_done."
        ),
        "condiciones": [
            "Se completan todos los campos con opciones válidas.",
            "Se envía el formulario con el nombre del negocio vacío.",
            "Se fuerza (vía API) un sector fuera del conjunto permitido.",
            "Un usuario ya con onboarding hecho intenta volver a la pantalla de onboarding.",
        ],
        "resultados": [
            "El sistema guarda el perfil, marca el onboarding como hecho y redirige al panel principal.",
            "El sistema muestra la validación de campo obligatorio y no envía la solicitud.",
            "El sistema responde 400 «Sector inválido» y no persiste el perfil.",
            "El sistema redirige al panel principal; no vuelve a exigir el onboarding.",
        ],
        "metodo": (
            "Partición de equivalencias: valores dentro del conjunto permitido vs fuera de "
            "él, para cada selector, y campo obligatorio vacío vs con contenido."
        ),
        "modulos_sistema": "auth.py (opciones_perfil, guardar_perfil), OnboardingPage.tsx.",
        "procedimiento": [
            "Ingresar por primera vez con una cuenta no administradora.",
            "Completar el formulario según cada condición.",
            "Enviar y observar la validación o la redirección.",
            "Registrar el resultado y capturar la pantalla.",
        ],
        "dependencias": ["Cuenta no administradora sin onboarding previo.", "Backend en ejecución."],
    },
    {
        "num": 4,
        "id_hu": "HU-1.4: Gestión de roles",
        "nombre": "Creación de Roles con Permisos",
        "modulo": "Rol",
        "tecnica": "Partición de equivalencias",
        "descripcion": (
            "Verificar que un administrador pueda crear roles con un conjunto de permisos "
            "válidos y que el sistema rechace nombres duplicados o permisos desconocidos."
        ),
        "entorno": "Editor de roles dentro de la sección de administración; requiere permiso de administración.",
        "parametros": ["Campo: Nombre del rol", "Campo: Descripción", "Selección múltiple: Permisos"],
        "respuesta_modulos": "GET /permissions provee el catálogo; POST /roles valida los permisos y crea el rol.",
        "condiciones": [
            "Se crea un rol con nombre único y permisos válidos del catálogo.",
            "Se crea un rol con un nombre que ya existe.",
            "Se envía un permiso que no está en el catálogo.",
            "Un usuario sin permiso de administración intenta crear un rol.",
        ],
        "resultados": [
            "El sistema crea el rol (201) y lo muestra en el listado.",
            "El sistema responde 409 «Ya existe un rol con ese nombre».",
            "El sistema responde 400 «Permisos desconocidos» y no crea el rol.",
            "El sistema responde 403 y no expone la operación.",
        ],
        "metodo": "Partición de equivalencias: nombre único vs duplicado; permisos válidos vs desconocidos; con vs sin permiso de administración.",
        "modulos_sistema": "auth.py (crear_rol, listar_permisos), UsersPage.tsx (editor de roles).",
        "procedimiento": [
            "Ingresar como administrador y abrir el editor de roles.",
            "Crear un rol por cada condición.",
            "Observar la respuesta del sistema y el listado.",
            "Registrar el resultado y capturar la pantalla.",
        ],
        "dependencias": ["Cuenta con permiso de administración.", "Catálogo de permisos disponible."],
    },
    {
        "num": 5,
        "id_hu": "HU-1.5: Editar rol",
        "nombre": "Editar Rol (Administrador Protegido)",
        "modulo": "Rol",
        "tecnica": "Tabla de decisión",
        "descripcion": (
            "Verificar que se pueda editar la descripción y los permisos de un rol, salvo "
            "los permisos del rol administrador, que están protegidos."
        ),
        "entorno": "Editor de roles; el rol administrador aparece marcado como protegido.",
        "parametros": ["Selección: Rol a editar", "Campo: Descripción", "Selección múltiple: Permisos"],
        "respuesta_modulos": "PATCH /roles/{id} aplica los cambios salvo cuando el rol es el administrador y se intentan cambiar sus permisos.",
        "condiciones": [
            "Se edita la descripción de un rol normal.",
            "Se cambian los permisos de un rol normal por un conjunto válido.",
            "Se intentan modificar los permisos del rol administrador.",
            "Se edita un rol con un id inexistente.",
        ],
        "resultados": [
            "El sistema actualiza la descripción y refleja el cambio.",
            "El sistema actualiza los permisos del rol.",
            "El sistema responde 400 «No se pueden modificar los permisos del rol administrador».",
            "El sistema responde 404 «El rol no existe».",
        ],
        "metodo": "Tabla de decisión: (rol normal / rol admin) × (edita descripción / edita permisos) → acción permitida o bloqueada.",
        "modulos_sistema": "auth.py (actualizar_rol), UsersPage.tsx.",
        "procedimiento": [
            "Ingresar como administrador y abrir el editor de roles.",
            "Aplicar cada edición según la condición.",
            "Observar la respuesta del sistema.",
            "Registrar el resultado y capturar la pantalla.",
        ],
        "dependencias": ["Deben existir un rol normal y el rol administrador.", "Cuenta con permiso de administración."],
    },
    {
        "num": 6,
        "id_hu": "HU-1.6: Eliminar rol",
        "nombre": "Eliminar Rol con Usuarios Asignados",
        "modulo": "Rol",
        "tecnica": "Tabla de decisión",
        "descripcion": "Verificar que solo se puedan eliminar roles sin usuarios asignados y que el rol administrador esté protegido.",
        "entorno": "Listado de roles con la acción Eliminar por fila.",
        "parametros": ["Selección: Rol a eliminar", "Botón: Eliminar"],
        "respuesta_modulos": "DELETE /roles/{id} verifica que el rol no sea el administrador y que no tenga usuarios asignados.",
        "condiciones": [
            "Se elimina un rol normal sin usuarios asignados.",
            "Se intenta eliminar un rol que tiene usuarios asignados.",
            "Se intenta eliminar el rol administrador.",
            "Se elimina un rol con un id inexistente.",
        ],
        "resultados": [
            "El sistema elimina el rol (204) y lo quita del listado.",
            "El sistema responde 400 «No se puede eliminar un rol con usuarios asignados».",
            "El sistema responde 400 «No se puede eliminar el rol administrador».",
            "El sistema responde 404 «El rol no existe».",
        ],
        "metodo": "Tabla de decisión: (rol admin / con usuarios / vacío / inexistente) → resultado esperado.",
        "modulos_sistema": "auth.py (eliminar_rol), UsersPage.tsx.",
        "procedimiento": [
            "Ingresar como administrador y abrir el listado de roles.",
            "Ejecutar Eliminar por cada condición.",
            "Observar la respuesta del sistema.",
            "Registrar el resultado y capturar la pantalla.",
        ],
        "dependencias": ["Un rol con usuarios asignados y otro vacío.", "Cuenta con permiso de administración."],
    },
    {
        "num": 7,
        "id_hu": "HU-1.7: Crear usuario",
        "nombre": "Creación de Cuenta de Usuario",
        "modulo": "Usuario",
        "tecnica": "Partición de equivalencias",
        "descripcion": "Verificar que un administrador cree cuentas asignando un rol existente y que el sistema rechace ids duplicados o roles inexistentes.",
        "entorno": "Formulario de alta de usuario en la sección de administración.",
        "parametros": ["Campo: Id de usuario", "Campo: Contraseña", "Selección: Rol", "Campo: Correo (opcional)"],
        "respuesta_modulos": "POST /users valida que el rol exista y que el id no esté en uso antes de crear la cuenta.",
        "condiciones": [
            "Se crea una cuenta con id único, contraseña válida y un rol existente.",
            "Se crea una cuenta con un id que ya existe.",
            "Se asigna un rol con id inexistente.",
            "Se deja vacío el id o la contraseña.",
        ],
        "resultados": [
            "El sistema crea la cuenta (201) y la muestra en el listado.",
            "El sistema responde 409 «Ya existe un usuario con ese id».",
            "El sistema responde 404 «El rol no existe».",
            "El sistema muestra la validación de campo obligatorio y no envía la solicitud.",
        ],
        "metodo": "Partición de equivalencias: id único vs duplicado; rol existente vs inexistente; campos completos vs vacíos.",
        "modulos_sistema": "auth.py (crear_usuario, listar_usuarios), UsersPage.tsx.",
        "procedimiento": [
            "Ingresar como administrador y abrir el alta de usuarios.",
            "Crear una cuenta por cada condición.",
            "Observar la respuesta del sistema y el listado.",
            "Registrar el resultado y capturar la pantalla.",
        ],
        "dependencias": ["Al menos un rol existente.", "Cuenta con permiso de administración."],
    },
    {
        "num": 8,
        "id_hu": "HU-1.8: Editar usuario",
        "nombre": "Editar Usuario (Rol, Estado y Contraseña)",
        "modulo": "Usuario",
        "tecnica": "Tabla de decisión",
        "descripcion": "Verificar la edición de rol, estado (activo/inactivo) y contraseña de una cuenta, y el efecto de desactivar una cuenta sobre su acceso.",
        "entorno": "Ficha de edición de usuario en la sección de administración.",
        "parametros": ["Selección: Rol", "Interruptor: Activo / Inactivo", "Campo: Nueva contraseña", "Campo: Correo"],
        "respuesta_modulos": "PATCH /users/{id} aplica los cambios; una cuenta inactiva queda rechazada en el login.",
        "condiciones": [
            "Se cambia el rol de una cuenta a otro rol existente.",
            "Se desactiva una cuenta activa (is_active = false).",
            "La cuenta desactivada intenta iniciar sesión.",
            "Se edita un usuario con un id inexistente.",
        ],
        "resultados": [
            "El sistema actualiza el rol y refleja los nuevos permisos.",
            "El sistema marca la cuenta como inactiva.",
            "El sistema rechaza el acceso con el mensaje genérico de credenciales inválidas.",
            "El sistema responde 404 «No existe el usuario».",
        ],
        "metodo": "Tabla de decisión: (cambia rol / desactiva / login tras desactivar / id inexistente) → resultado esperado.",
        "modulos_sistema": "auth.py (actualizar_usuario, login), UsersPage.tsx.",
        "procedimiento": [
            "Ingresar como administrador y abrir la edición de una cuenta.",
            "Aplicar cada cambio según la condición.",
            "Para la condición 3: cerrar sesión e intentar entrar con la cuenta desactivada.",
            "Registrar el resultado y capturar la pantalla.",
        ],
        "dependencias": ["Una cuenta editable distinta del administrador.", "Cuenta con permiso de administración."],
    },
    {
        "num": 9,
        "id_hu": "HU-1.9: Control de acceso por rol",
        "nombre": "Visibilidad de Secciones según Permiso",
        "modulo": "Permisos",
        "tecnica": "Tabla de decisión",
        "descripcion": "Verificar que el menú y las rutas muestren únicamente las secciones para las que el rol tiene permiso.",
        "entorno": "Panel principal de SPC; el menú lateral se arma a partir de los permisos del rol.",
        "parametros": ["Rol del usuario autenticado", "Secciones: Inicio, Ventas, Compras, Inventario, Usuarios, Acerca de"],
        "respuesta_modulos": "GET /auth/me entrega los permisos; useSeccionesVisibles filtra las secciones renderizadas.",
        "condiciones": [
            "Un rol con permiso solo de Ventas inicia sesión.",
            "Ese rol intenta navegar por URL a la sección Usuarios (sin permiso).",
            "Un rol sin ninguna sección permitida inicia sesión.",
            "El administrador inicia sesión.",
        ],
        "resultados": [
            "El menú muestra únicamente Ventas; las demás secciones no aparecen.",
            "El sistema redirige a la primera sección permitida; no muestra Usuarios.",
            "El sistema muestra el aviso «Su rol no tiene acceso a ninguna sección».",
            "El menú muestra todas las secciones, incluida Usuarios.",
        ],
        "metodo": "Tabla de decisión: permisos del rol × sección solicitada → visible / redirigido / bloqueado.",
        "modulos_sistema": "auth.py (me, listar_permisos), useSeccionesVisibles.ts, App.tsx, Layout.tsx.",
        "procedimiento": [
            "Preparar roles con distintos permisos.",
            "Iniciar sesión con cada rol y revisar el menú.",
            "Intentar la navegación directa por URL a una sección no permitida.",
            "Registrar el resultado y capturar la pantalla.",
        ],
        "dependencias": ["Roles con permisos distintos y el rol administrador.", "Backend en ejecución."],
    },
    {
        "num": 10,
        "id_hu": "HU-2.1: Análisis de ventas",
        "nombre": "Análisis 3×3 de VENTAS (Regresión + Clasificación + Clustering)",
        "modulo": "Ventas",
        "tecnica": "Caso de uso",
        "descripcion": (
            "Verificar que, dado un conjunto de datos de ventas, el sistema entrene y "
            "devuelva los tres análisis (regresión, clasificación y clustering) con métricas honestas."
        ),
        "entorno": "Sección Ventas del panel; carga de datos o uso del ejemplo integrado.",
        "parametros": ["Serie temporal de ventas (fecha, serie, valor)", "Características (features) declaradas", "Horizonte de predicción"],
        "respuesta_modulos": "POST /v2/ventas ejecuta el pipeline 3×3 y devuelve predicción, clases y agrupaciones con sus métricas.",
        "condiciones": [
            "Se envía un dataset de ventas válido con suficiente historia.",
            "Se ejecuta la demo con datos sintéticos (GET /v2/ventas/demo).",
            "Se envía un dataset con muy pocos registros para entrenar.",
            "Se agregan más características informativas al mismo dataset.",
        ],
        "resultados": [
            "El sistema devuelve la regresión (WAPE dentro del rango esperado ~0,05–0,16), las clases y el clustering, con sus gráficos.",
            "El sistema devuelve un resultado consistente sobre los datos de ejemplo.",
            "El sistema responde con un error controlado indicando datos insuficientes, sin caerse.",
            "La métrica no empeora respecto al piso de ruido; puede mejorar o mantenerse.",
        ],
        "metodo": "Caso de uso: recorrido del flujo real de análisis de ventas, contrastando la métrica contra el piso de ruido y el baseline naive.",
        "modulos_sistema": "dominios_3x3.py (/ventas, /ventas/demo), núcleo ML (nucleo.py), VistaV3Modulo.tsx, SalesPage.tsx.",
        "procedimiento": [
            "Abrir la sección Ventas.",
            "Cargar el dataset o ejecutar la demo.",
            "Lanzar el análisis y revisar las tres salidas y sus métricas.",
            "Registrar el resultado y capturar la pantalla.",
        ],
        "dependencias": ["Dataset de ventas con formato del dominio.", "Backend y modelos disponibles."],
    },
    {
        "num": 11,
        "id_hu": "HU-2.2: Análisis de compras",
        "nombre": "Análisis 3×3 de COMPRAS",
        "modulo": "Compras",
        "tecnica": "Caso de uso",
        "descripcion": "Verificar el pipeline 3×3 aplicado al dominio de compras y la coherencia de sus tres salidas.",
        "entorno": "Sección Compras del panel; carga de datos o demo integrada.",
        "parametros": ["Serie temporal de compras", "Características declaradas", "Horizonte de predicción"],
        "respuesta_modulos": "POST /v2/compras ejecuta el pipeline 3×3 del dominio compras.",
        "condiciones": [
            "Se envía un dataset de compras válido.",
            "Se ejecuta la demo de compras (GET /v2/compras/demo).",
            "Se envía un dataset con columnas que no corresponden al esquema del dominio.",
        ],
        "resultados": [
            "El sistema devuelve regresión, clasificación y clustering con métricas dentro del rango esperado.",
            "El sistema devuelve un resultado consistente sobre el ejemplo.",
            "El sistema responde con un error de validación claro, sin procesar datos inválidos.",
        ],
        "metodo": "Caso de uso: flujo de análisis de compras, verificando métricas honestas (sin fuga entre train/valid/test).",
        "modulos_sistema": "dominios_3x3.py (/compras, /compras/demo), núcleo ML, PurchasesPage.tsx.",
        "procedimiento": [
            "Abrir la sección Compras.",
            "Cargar el dataset o ejecutar la demo.",
            "Lanzar el análisis y revisar las salidas.",
            "Registrar el resultado y capturar la pantalla.",
        ],
        "dependencias": ["Dataset de compras con formato del dominio.", "Backend y modelos disponibles."],
    },
    {
        "num": 12,
        "id_hu": "HU-2.3: Análisis de almacén",
        "nombre": "Análisis 3×3 de ALMACÉN",
        "modulo": "Almacén",
        "tecnica": "Caso de uso",
        "descripcion": "Verificar el pipeline 3×3 aplicado al dominio de almacén / inventario y la coherencia de sus salidas.",
        "entorno": "Sección Inventario del panel; carga de datos o demo integrada.",
        "parametros": ["Serie temporal de almacén", "Características declaradas", "Horizonte de predicción"],
        "respuesta_modulos": "POST /v2/almacen ejecuta el pipeline 3×3 del dominio almacén.",
        "condiciones": [
            "Se envía un dataset de almacén válido.",
            "Se ejecuta la demo de almacén (GET /v2/almacen/demo).",
            "Se envía un dataset con valores faltantes en la serie.",
        ],
        "resultados": [
            "El sistema devuelve las tres salidas con métricas dentro del rango esperado.",
            "El sistema devuelve un resultado consistente sobre el ejemplo.",
            "El sistema maneja los faltantes (imputación o aviso) sin caerse.",
        ],
        "metodo": "Caso de uso: flujo de análisis de almacén, verificando el manejo de datos incompletos y la honestidad de las métricas.",
        "modulos_sistema": "dominios_3x3.py (/almacen, /almacen/demo), núcleo ML, InventoryPage.tsx.",
        "procedimiento": [
            "Abrir la sección Inventario.",
            "Cargar el dataset o ejecutar la demo.",
            "Lanzar el análisis y revisar las salidas.",
            "Registrar el resultado y capturar la pantalla.",
        ],
        "dependencias": ["Dataset de almacén con formato del dominio.", "Backend y modelos disponibles."],
    },
    {
        "num": 13,
        "id_hu": "HU-2.4: Catálogo de consultas",
        "nombre": "Ejecutar Consulta del Catálogo (JSON)",
        "modulo": "Catálogo",
        "tecnica": "Partición de equivalencias",
        "descripcion": "Verificar que el usuario ejecute una de las consultas predefinidas del catálogo enviando datos en JSON y obtenga el análisis correspondiente.",
        "entorno": "Sección de reportes / catálogo con las 30 consultas predefinidas por módulo.",
        "parametros": ["Selección: Módulo", "Selección: Consulta del catálogo", "Cuerpo JSON con los datos"],
        "respuesta_modulos": "GET /catalogo lista las consultas; POST /{modulo} ejecuta la consulta con el JSON enviado.",
        "condiciones": [
            "Se ejecuta una consulta existente con un JSON bien formado.",
            "Se envía un JSON con un campo obligatorio ausente.",
            "Se selecciona un módulo que no existe en el catálogo.",
            "Se ejecuta la demo del módulo (GET /{modulo}/demo).",
        ],
        "resultados": [
            "El sistema devuelve el análisis solicitado y lo muestra en tabla y gráficos.",
            "El sistema responde con un error de validación indicando el campo faltante.",
            "El sistema responde con un error de módulo desconocido.",
            "El sistema devuelve un resultado de ejemplo consistente.",
        ],
        "metodo": "Partición de equivalencias: JSON válido vs inválido; módulo existente vs inexistente; consulta con datos vs demo.",
        "modulos_sistema": "catalogo_v3.py (catálogo, POST /{modulo}, demo), CatalogTable.tsx, ResultTable.tsx.",
        "procedimiento": [
            "Abrir el catálogo y elegir módulo y consulta.",
            "Enviar el JSON por cada condición.",
            "Observar el resultado o el error.",
            "Registrar el resultado y capturar la pantalla.",
        ],
        "dependencias": ["Catálogo cargado.", "Backend en ejecución."],
    },
    {
        "num": 14,
        "id_hu": "HU-2.5: Carga de archivo",
        "nombre": "Carga de Archivo Excel/JSON y Validación",
        "modulo": "Carga",
        "tecnica": "Partición de equivalencias y análisis de valores borde",
        "descripcion": "Verificar la carga de datos por archivo (Excel o JSON), la validación previa del contenido y el manejo de archivos mal formados.",
        "entorno": "Componente de carga con validación (ValidacionCarga) en las secciones de análisis.",
        "parametros": ["Archivo: .xlsx o .json", "Columnas requeridas del dominio", "Tamaño / número de filas"],
        "respuesta_modulos": "POST /{modulo}/archivo procesa el archivo, valida el esquema y devuelve el análisis o los errores de validación.",
        "condiciones": [
            "Se carga un Excel válido con las columnas requeridas.",
            "Se carga un JSON válido equivalente.",
            "Se carga un archivo con columnas faltantes o mal nombradas.",
            "Se carga un archivo de formato no soportado (por ejemplo .csv o .txt).",
            "Se carga un archivo vacío o con cero filas de datos.",
        ],
        "resultados": [
            "El sistema valida, procesa el archivo y muestra el análisis.",
            "El sistema procesa el JSON de la misma forma que el Excel.",
            "El sistema muestra la validación con las columnas que faltan y no procesa.",
            "El sistema rechaza el archivo indicando el formato no soportado.",
            "El sistema avisa que el archivo no contiene datos y no procesa.",
        ],
        "metodo": "Partición de equivalencias: formato soportado vs no soportado; esquema completo vs incompleto. Valores borde: archivo vacío.",
        "modulos_sistema": "catalogo_v3.py (POST /{modulo}/archivo), ValidacionCarga.tsx.",
        "procedimiento": [
            "Abrir una sección de análisis y elegir cargar archivo.",
            "Cargar un archivo por cada condición.",
            "Observar la validación y el resultado.",
            "Registrar el resultado y capturar la pantalla.",
        ],
        "dependencias": ["Archivos de prueba (válidos e inválidos).", "Backend en ejecución."],
    },
    {
        "num": 15,
        "id_hu": "HU-2.6: Historial de análisis",
        "nombre": "Historial y Detalle de Análisis",
        "modulo": "Historial",
        "tecnica": "Caso de uso",
        "descripcion": "Verificar que el sistema registre cada análisis, permita listar el historial por módulo y abrir el detalle de una ejecución.",
        "entorno": "Sección de historial con el listado de análisis previos y su detalle.",
        "parametros": ["Selección: Módulo", "Identificador de la predicción (prediction_id)"],
        "respuesta_modulos": "GET /{modulo}/historial lista las ejecuciones; GET /historial/{prediction_id} devuelve el detalle.",
        "condiciones": [
            "Se consulta el historial de un módulo con análisis previos.",
            "Se abre el detalle de una predicción existente.",
            "Se consulta el detalle de un prediction_id inexistente.",
            "Se consulta el historial de un módulo sin análisis previos.",
        ],
        "resultados": [
            "El sistema lista las ejecuciones con su fecha y métricas resumidas.",
            "El sistema muestra el detalle completo (datos, resultado y métricas) de esa ejecución.",
            "El sistema responde con un error de recurso no encontrado.",
            "El sistema muestra un estado vacío indicando que aún no hay análisis.",
        ],
        "metodo": "Caso de uso: recorrido del historial y del detalle, verificando la persistencia de cada ejecución.",
        "modulos_sistema": "catalogo_v3.py (historial, detalle), HistorialV3.tsx, HistoryPreview.tsx.",
        "procedimiento": [
            "Ejecutar algún análisis para poblar el historial.",
            "Abrir el historial del módulo y luego el detalle.",
            "Probar un prediction_id inexistente y un módulo sin historial.",
            "Registrar el resultado y capturar la pantalla.",
        ],
        "dependencias": ["Al menos un análisis registrado.", "Persistencia (PostgreSQL) conectada."],
    },
    {
        "num": 16,
        "id_hu": "HU-2.7: Servir modelo y reportes",
        "nombre": "Servir Modelo Guardado y Exportar Reportes",
        "modulo": "Reportes",
        "tecnica": "Caso de uso",
        "descripcion": (
            "Verificar que el sistema prediga usando el modelo adoptado (sin reentrenar), permita "
            "elegir la versión a servir y exporte los reportes y gráficos."
        ),
        "entorno": "Sección de reportes con la predicción sobre el modelo guardado y las opciones de exportación.",
        "parametros": ["Selección: Dominio (ventas/compras/almacén)", "Identificador del modelo / versión", "Formato de exportación (Excel, gráficos)"],
        "respuesta_modulos": "POST /v2/{dominio}/predecir usa el modelo adoptado; POST /v2/{dominio}/modelos/{id}/servir elige la versión.",
        "condiciones": [
            "Se predice con el modelo adoptado del dominio (sin reentrenar).",
            "Se cambia la versión servida a otra versión guardada.",
            "Se solicita predecir en un dominio sin ningún modelo adoptado.",
            "Se exporta el reporte a Excel y se descargan los gráficos.",
        ],
        "resultados": [
            "El sistema devuelve la predicción de forma inmediata reutilizando el modelo adoptado.",
            "El sistema pasa a servir la versión seleccionada y lo refleja en la auditoría.",
            "El sistema responde con un error indicando que no hay modelo adoptado.",
            "El sistema genera el Excel y los gráficos descargables con los datos del análisis.",
        ],
        "metodo": "Caso de uso: predicción servida desde el modelo adoptado, cambio de versión y exportación de resultados.",
        "modulos_sistema": "dominios_3x3.py (/{dominio}/predecir, /{dominio}/modelos/{id}/servir), SeccionReportes.tsx, ReporteCard.tsx, charts/*.",
        "procedimiento": [
            "Adoptar un modelo para el dominio.",
            "Ejecutar la predicción servida y cambiar la versión.",
            "Exportar el reporte a Excel y descargar los gráficos.",
            "Registrar el resultado y capturar la pantalla.",
        ],
        "dependencias": ["Un modelo entrenado y adoptado por dominio.", "Persistencia y bucket de modelos disponibles."],
    },
]


# ================================================================================
# Utilidades de estilo para el .docx
# ================================================================================
def _sombrear(celda, hex_fill: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    celda._tc.get_or_add_tcPr().append(shd)


def _texto(celda, texto: str, *, negrita=False, blanco=False, tam=10) -> None:
    celda.text = ""
    p = celda.paragraphs[0]
    lineas = str(texto).split("\n")
    for i, linea in enumerate(lineas):
        run = p.add_run(linea)
        run.font.size = Pt(tam)
        run.font.bold = negrita
        run.font.color.rgb = BLANCO if blanco else GRIS_TXT
        if i < len(lineas) - 1:
            run.add_break()


def _fila_seccion(tabla, titulo: str) -> None:
    fila = tabla.add_row()
    celda = fila.cells[0].merge(fila.cells[1])
    _texto(celda, titulo, negrita=True, blanco=True)
    _sombrear(celda, AZUL_MEDIO)


def _fila_kv(tabla, etiqueta: str, valor: str) -> None:
    fila = tabla.add_row()
    _texto(fila.cells[0], etiqueta, negrita=True)
    _sombrear(fila.cells[0], AZUL_CLARO)
    _texto(fila.cells[1], valor)


def _fila_titulo(tabla, titulo: str) -> None:
    fila = tabla.add_row()
    celda = fila.cells[0].merge(fila.cells[1])
    _texto(celda, titulo, negrita=True, blanco=True, tam=11)
    _sombrear(celda, AZUL_OSCURO)


def _viñeta(items: list[str]) -> str:
    return "\n".join(f"• {x}" for x in items)


def _numerado(items: list[str]) -> str:
    return "\n".join(f"{i}. {x}" for i, x in enumerate(items, 1))


# ================================================================================
# Construcción del .docx
# ================================================================================
def construir_docx(destino: Path) -> None:
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)

    # --- Portada ---------------------------------------------------------------
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("DOCUMENTO DE PRUEBAS DE CAJA NEGRA")
    r.bold = True
    r.font.size = Pt(20)
    r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rs = sub.add_run("SPC — Sistema de Predicción y Comercialización")
    rs.bold = True
    rs.font.size = Pt(14)

    for linea in (
        "Tipo de Prueba: Caja Negra",
        "Técnicas: Partición de Equivalencias, Análisis de Valores Borde, Tabla de Decisión, Caso de Uso",
        f"Total de Escenarios: {len(ESCENARIOS)}",
        "Universidad Privada Antenor Orrego — Proyecto de Taller Integrador I",
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(linea).font.size = Pt(11)

    doc.add_paragraph()

    # --- Tabla índice ----------------------------------------------------------
    h = doc.add_paragraph()
    hr = h.add_run("Índice de Escenarios de Prueba")
    hr.bold = True
    hr.font.size = Pt(13)

    idx = doc.add_table(rows=1, cols=5)
    idx.style = "Table Grid"
    idx.alignment = WD_TABLE_ALIGNMENT.CENTER
    cab = ["N°", "ID HU", "Escenario", "Módulo", "Técnica"]
    for celda, txt in zip(idx.rows[0].cells, cab):
        _texto(celda, txt, negrita=True, blanco=True)
        _sombrear(celda, AZUL_OSCURO)
    for e in ESCENARIOS:
        fila = idx.add_row().cells
        _texto(fila[0], str(e["num"]))
        _texto(fila[1], e["id_hu"].split(":")[0])
        _texto(fila[2], e["nombre"])
        _texto(fila[3], e["modulo"])
        _texto(fila[4], e["tecnica"])

    doc.add_page_break()

    # --- Un cuadro por escenario ----------------------------------------------
    for e in ESCENARIOS:
        tabla = doc.add_table(rows=0, cols=2)
        tabla.style = "Table Grid"
        tabla.alignment = WD_TABLE_ALIGNMENT.CENTER

        _fila_titulo(tabla, f"Escenario {e['num']}.0 — {e['nombre']}")
        _fila_kv(tabla, "Historia de Usuario", e["id_hu"])
        _fila_kv(tabla, "Módulo / Contexto", e["modulo"])
        _fila_kv(tabla, "Tipo de Prueba", "Prueba de Caja Negra")
        _fila_kv(tabla, "Técnica Aplicada", e["tecnica"])

        _fila_seccion(tabla, "DATOS DE ENTRADA")
        _fila_kv(tabla, "Descripción", e["descripcion"])
        _fila_kv(tabla, "Entorno", e["entorno"])
        _fila_kv(tabla, "Parámetros involucrados", _viñeta(e["parametros"]))
        _fila_kv(tabla, "Respuesta de otros módulos", e["respuesta_modulos"])

        _fila_seccion(tabla, "CONDICIONES INICIALES (CASOS DE PRUEBA)")
        for i, cond in enumerate(e["condiciones"], 1):
            _fila_kv(tabla, f"Condición {i}", cond)

        _fila_seccion(tabla, "DATOS DE SALIDA")
        for i, res in enumerate(e["resultados"], 1):
            _fila_kv(tabla, f"Resultado condición {i}", res)
        _fila_kv(tabla, "Estado final de variables", "Se adjuntan capturas de pantalla como evidencia.")
        for i in range(1, len(e["condiciones"]) + 1):
            _fila_kv(tabla, f"Evidencia condición {i}", "")

        _fila_seccion(tabla, "REQUISITOS DE CONFIGURACIÓN")
        _fila_kv(tabla, "Método de Prueba", e["metodo"])
        _fila_kv(tabla, "Módulos del sistema", e["modulos_sistema"])
        _fila_kv(tabla, "Hardware y Software", HARDWARE_SW)
        _fila_kv(tabla, "Procedimiento de ejecución", _numerado(e["procedimiento"]))
        _fila_kv(tabla, "Dependencias", _viñeta(e["dependencias"]))
        _fila_kv(tabla, "Estado de la prueba", ESTADO)
        _fila_kv(tabla, "Responsable", RESPONSABLE)

        doc.add_paragraph()
        if e["num"] != ESCENARIOS[-1]["num"]:
            doc.add_page_break()

    doc.save(str(destino))


# ================================================================================
# Construcción del .md fuente
# ================================================================================
def construir_md(destino: Path) -> None:
    L: list[str] = []
    L.append("# Documento de Pruebas de Caja Negra")
    L.append("")
    L.append("**SPC — Sistema de Predicción y Comercialización**")
    L.append("")
    L.append("- **Tipo de Prueba:** Caja Negra")
    L.append("- **Técnicas:** Partición de Equivalencias, Análisis de Valores Borde, Tabla de Decisión, Caso de Uso")
    L.append(f"- **Total de Escenarios:** {len(ESCENARIOS)}")
    L.append("- **Universidad Privada Antenor Orrego** — Proyecto de Taller Integrador I")
    L.append("")
    L.append("## Índice de Escenarios de Prueba")
    L.append("")
    L.append("| N° | ID HU | Escenario | Módulo | Técnica |")
    L.append("|----|-------|-----------|--------|---------|")
    for e in ESCENARIOS:
        L.append(f"| {e['num']} | {e['id_hu'].split(':')[0]} | {e['nombre']} | {e['modulo']} | {e['tecnica']} |")
    L.append("")

    for e in ESCENARIOS:
        L.append(f"## Escenario {e['num']}.0 — {e['nombre']}")
        L.append("")
        L.append(f"- **Historia de Usuario:** {e['id_hu']}")
        L.append(f"- **Módulo / Contexto:** {e['modulo']}")
        L.append("- **Tipo de Prueba:** Prueba de Caja Negra")
        L.append(f"- **Técnica Aplicada:** {e['tecnica']}")
        L.append("")
        L.append("### Datos de entrada")
        L.append("")
        L.append(f"- **Descripción:** {e['descripcion']}")
        L.append(f"- **Entorno:** {e['entorno']}")
        L.append("- **Parámetros involucrados:**")
        for p in e["parametros"]:
            L.append(f"  - {p}")
        L.append(f"- **Respuesta de otros módulos:** {e['respuesta_modulos']}")
        L.append("")
        L.append("### Condiciones iniciales (casos de prueba)")
        L.append("")
        L.append("| # | Condición | Resultado esperado |")
        L.append("|---|-----------|--------------------|")
        for i, (c, r) in enumerate(zip(e["condiciones"], e["resultados"]), 1):
            L.append(f"| {i} | {c} | {r} |")
        L.append("")
        L.append("_Estado final de variables: se adjuntan capturas de pantalla como evidencia._")
        L.append("")
        L.append("### Requisitos de configuración")
        L.append("")
        L.append(f"- **Método de Prueba:** {e['metodo']}")
        L.append(f"- **Módulos del sistema:** {e['modulos_sistema']}")
        L.append(f"- **Hardware y Software:** {HARDWARE_SW}")
        L.append("- **Procedimiento de ejecución:**")
        for i, s in enumerate(e["procedimiento"], 1):
            L.append(f"  {i}. {s}")
        L.append("- **Dependencias:**")
        for d in e["dependencias"]:
            L.append(f"  - {d}")
        L.append(f"- **Estado de la prueba:** {ESTADO}")
        L.append(f"- **Responsable:** {RESPONSABLE}")
        L.append("")

    destino.write_text("\n".join(L), encoding="utf-8")


def main() -> None:
    final = Path(__file__).resolve().parent.parent / "final"
    docx_path = final / "Pruebas de Caja Negra SPC.docx"
    md_path = final / "Pruebas de Caja Negra SPC.md"
    construir_docx(docx_path)
    construir_md(md_path)
    print(f"OK  ->  {docx_path}")
    print(f"OK  ->  {md_path}")


if __name__ == "__main__":
    main()
