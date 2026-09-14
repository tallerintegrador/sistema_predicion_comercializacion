"""Genera el Documento de Pruebas End to End (E2E) del sistema SPC.

Las pruebas E2E se ejecutan con **Selenium WebDriver sobre Chrome real**: la suite
``tests/e2e_selenium/`` levanta el backend (uvicorn, base SQLite temporal) y el frontend
(Vite), abre el navegador y recorre la aplicación como una persona — incluida la subida de
los Excel del negocio de ejemplo. Cada corrida deja video, capturas y logs en
``final/evidencias/pruebas/e2e_selenium/``.

La lista ``FLUJOS`` es la única fuente de verdad; emite ``.docx`` + ``.md`` con el estilo de
``_docx_estilo`` e incrusta una captura real de la última corrida por flujo.

Uso:
    python scripts/gen_pruebas_e2e.py

Salida:
    final/Pruebas End to End SPC.docx
    final/Pruebas End to End SPC.md
"""

from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import _docx_estilo as est  # noqa: E402

RAIZ = Path(__file__).resolve().parent.parent
EVIDENCIA = RAIZ / "final" / "evidencias" / "pruebas" / "e2e_selenium"
CAPTURAS = EVIDENCIA / "capturas"

ENTORNO = (
    "Selenium WebDriver 4 sobre Google Chrome real (ventana visible, resolución 1600×1000). "
    "La propia suite levanta el backend (uvicorn, spc.api.main:app) y el frontend (Vite) en "
    "puertos libres; el control de acceso está activo y la persistencia usa una base SQLite "
    "temporal por corrida (no toca la base real ni Supabase). Datos: los Excel del negocio de "
    "ejemplo examples/datos_excel_pyme/botica_farmasalud/."
)
COMANDO = "venv/Scripts/python -m pytest tests/e2e_selenium -m selenium -v"
EVIDENCIA_TXT = "final/evidencias/pruebas/e2e_selenium/ (video, capturas y logs)"

# ================================================================================
# Flujos E2E (única fuente de verdad; archivos reales de tests/e2e_selenium/)
# ================================================================================
FLUJOS: list[dict] = [
    {
        "num": 1,
        "archivo": "tests/e2e_selenium/test_01_login.py",
        "flujo": "Ingreso al sistema y cierre de sesión",
        "pantallas": "Login → Panel principal → Login",
        "precondicion": "Aplicación abierta en el navegador; cuentas de demostración sembradas.",
        "pasos": [
            "Abrir la aplicación: se muestra la pantalla de ingreso.",
            "Ingresar con contraseña incorrecta.",
            "Ingresar con las credenciales del administrador (256370).",
            "Verificar a qué servidor apunta el panel y cerrar sesión.",
        ],
        "verificacion": (
            "El intento fallido muestra un mensaje genérico y no deja pasar; el ingreso correcto "
            "abre el panel con el usuario en la barra lateral; el cierre de sesión vuelve al login; "
            "el log del backend registra POST /auth/login."
        ),
        "captura": "003_",
        "resultado": "PASA",
    },
    {
        "num": 2,
        "archivo": "tests/e2e_selenium/test_02_recuperar_password.py",
        "flujo": "Solicitud de restablecimiento de contraseña",
        "pantallas": "Login → ¿Olvidaste tu contraseña?",
        "precondicion": "Sin sesión iniciada.",
        "pasos": [
            "Abrir «¿Olvidaste tu contraseña?».",
            "Enviar un correo que no existe en el sistema.",
            "Leer la respuesta y volver al login.",
        ],
        "verificacion": (
            "La respuesta es genérica («si la cuenta existe…»): no revela si el correo está "
            "registrado; el backend registra la solicitud (POST /auth/forgot)."
        ),
        "captura": "007_",
        "resultado": "PASA",
    },
    {
        "num": 3,
        "archivo": "tests/e2e_selenium/test_03_ventas.py",
        "flujo": "Ventas: subir el Excel del negocio y obtener los 10 análisis",
        "pantallas": "Panel → Ventas (3 pasos guiados)",
        "precondicion": "Sesión de administrador; archivo ventas.xlsx del negocio de ejemplo.",
        "pasos": [
            "Entrar al módulo Ventas.",
            "Subir ventas.xlsx en el Paso 3 (el navegador adjunta el archivo real).",
            "Esperar a que el sistema entrene y devuelva los resultados.",
            "Abrir el detalle técnico de un reporte.",
        ],
        "verificacion": (
            "Llegan 10 reportes con los tres tipos (predicción, alerta y segmento); la carga "
            "informa las columnas reconocidas; el detalle técnico muestra modelo y métrica; el "
            "backend registra POST /v3/ventas/archivo."
        ),
        "captura": "011_",
        "resultado": "PASA",
    },
    {
        "num": 4,
        "archivo": "tests/e2e_selenium/test_04_compras.py",
        "flujo": "Compras: análisis de abastecimiento desde el Excel",
        "pantallas": "Panel → Compras",
        "precondicion": "Sesión de administrador; archivo compras.xlsx.",
        "pasos": [
            "Entrar al módulo Compras.",
            "Subir compras.xlsx.",
            "Revisar las preguntas que responde cada tarjeta.",
        ],
        "verificacion": (
            "Se generan los 10 análisis del módulo (cantidad a pedir, tiempo de entrega, costo y "
            "cumplimiento del proveedor) y cada tarjeta indica la pregunta que responde."
        ),
        "captura": "016_",
        "resultado": "PASA",
    },
    {
        "num": 5,
        "archivo": "tests/e2e_selenium/test_05_almacen.py",
        "flujo": "Almacén: reposición y clasificación ABC desde el Excel",
        "pantallas": "Panel → Almacén",
        "precondicion": "Sesión de administrador; archivo almacen.xlsx.",
        "pasos": [
            "Entrar al módulo Almacén.",
            "Subir almacen.xlsx.",
            "Comprobar que hay segmentaciones entre los resultados.",
        ],
        "verificacion": (
            "Se generan los 10 análisis, incluida al menos una segmentación (clasificación ABC / "
            "agrupación de productos); el backend registra POST /v3/almacen/archivo."
        ),
        "captura": "019_",
        "resultado": "PASA",
    },
    {
        "num": 6,
        "archivo": "tests/e2e_selenium/test_06_historial.py",
        "flujo": "Historial: los análisis quedan guardados y se pueden reabrir",
        "pantallas": "Ventas → Paso 4 (Historial de análisis)",
        "precondicion": "Haber ejecutado los análisis de los flujos 3 a 5.",
        "pasos": [
            "Volver al módulo Ventas.",
            "Revisar la tabla del historial.",
            "Pulsar «Ver resultados» de un análisis anterior.",
        ],
        "verificacion": (
            "El historial lista los análisis con fecha, filas y número de reportes; al abrirlo se "
            "muestran los valores predichos guardados (GET /v3/ventas/historial)."
        ),
        "captura": "022_",
        "resultado": "PASA",
    },
    {
        "num": 7,
        "archivo": "tests/e2e_selenium/test_07_usuarios_roles.py",
        "flujo": "Administración: crear un rol con permisos y un usuario",
        "pantallas": "Panel → Usuarios y permisos",
        "precondicion": "Sesión de administrador (único rol con la acción de administrar).",
        "pasos": [
            "Entrar a «Usuarios y permisos».",
            "Crear el rol e2e_rol con los tres módulos y la acción de predecir.",
            "Crear el usuario e2e_test con ese rol y un correo de contacto.",
            "Comprobar las tablas de roles y usuarios.",
        ],
        "verificacion": (
            "El rol y el usuario quedan persistidos y aparecen en las tablas; el usuario figura "
            "como Activo y con su correo; el backend registra POST /roles y POST /users."
        ),
        "captura": "027_",
        "resultado": "PASA",
    },
    {
        "num": 8,
        "archivo": "tests/e2e_selenium/test_08_rbac_usuario_nuevo.py",
        "flujo": "Usuario nuevo: onboarding, permisos aplicados y datos aislados",
        "pantallas": "Login → Onboarding → Panel (menú recortado) → Ventas",
        "precondicion": "El usuario e2e_test creado en el flujo 7.",
        "pasos": [
            "Ingresar como e2e_test (primer ingreso).",
            "Completar el onboarding del negocio (nombre, sector, tamaño, región y moneda).",
            "Revisar el menú y escribir a mano la ruta /users.",
            "Abrir Ventas, revisar el historial y generar su propio pronóstico.",
        ],
        "verificacion": (
            "El menú no muestra «Usuarios y permisos» y la ruta escrita a mano redirige (el "
            "permiso se aplica de verdad); el historial de la cuenta nueva arranca vacío — los "
            "datos del administrador no se filtran — y su análisis devuelve los 10 reportes."
        ),
        "captura": "030_",
        "resultado": "PASA",
    },
    {
        "num": 9,
        "archivo": "tests/e2e_selenium/test_09_reset_password_completo.py",
        "flujo": "Restablecer la contraseña de punta a punta con el enlace",
        "pantallas": "Login → ¿Olvidaste tu contraseña? → /reset?token=… → Login",
        "precondicion": "Cuenta con correo (e2e_test); SMTP en modo desarrollo (enlace al log).",
        "pasos": [
            "Solicitar el enlace desde el login con el correo de la cuenta.",
            "Tomar el enlace emitido por el servidor y abrirlo en el navegador.",
            "Fijar la contraseña nueva y confirmarla.",
            "Reutilizar el mismo enlace y volver a ingresar con la contraseña nueva.",
        ],
        "verificacion": (
            "El enlace permite fijar la contraseña una sola vez: el segundo intento se rechaza; "
            "la cuenta entra con la contraseña nueva."
        ),
        "captura": "036_",
        "resultado": "PASA",
    },
    {
        "num": 10,
        "archivo": "tests/e2e_selenium/test_10_navegacion.py",
        "flujo": "Recorrido por todas las secciones y «Acerca del sistema»",
        "pantallas": "Inicio, Ventas, Compras, Almacén, Usuarios, Acerca del sistema",
        "precondicion": "Sesión de administrador.",
        "pasos": [
            "Recorrer una por una las secciones del menú.",
            "Comprobar que cada sección carga contenido.",
            "Abrir «Acerca del sistema».",
            "Revisar la consola del navegador.",
        ],
        "verificacion": (
            "Ninguna sección queda en blanco; «Acerca del sistema» explica el funcionamiento y la "
            "exactitud; la consola no acumula errores de la aplicación."
        ),
        "captura": "039_",
        "resultado": "PASA",
    },
    {
        "num": 11,
        "archivo": "tests/e2e_selenium/test_11_errores_carga.py",
        "flujo": "Errores de carga explicados, sin caídas",
        "pantallas": "Ventas → Paso 3 (subida de archivo)",
        "precondicion": "Sesión de administrador; un .txt y un .xlsx corrupto.",
        "pasos": [
            "Subir un archivo .txt (extensión no admitida).",
            "Subir un .xlsx corrupto.",
            "Seguir usando la aplicación.",
        ],
        "verificacion": (
            "El .txt lo rechaza la interfaz con un aviso claro; el Excel corrupto devuelve un error "
            "controlado (422, no 500) que se muestra en el panel de error; la aplicación sigue "
            "operativa, sin pantalla en blanco."
        ),
        "captura": "049_",
        "resultado": "PASA",
    },
]


def _captura(prefijo: str) -> Path | None:
    """Ruta de la captura de la última corrida cuyo nombre empieza por el prefijo dado."""
    if not CAPTURAS.is_dir():
        return None
    for archivo in sorted(CAPTURAS.glob(f"{prefijo}*.png")):
        return archivo
    return None


def construir_docx(destino: Path) -> None:
    from docx.shared import Inches

    doc = est.nuevo_documento()
    est.portada(
        doc,
        "DOCUMENTO DE PRUEBAS END TO END (E2E)",
        [
            "Tipo de Prueba: End to End sobre la aplicación completa (navegador real)",
            "Herramienta: Selenium WebDriver 4 + Chrome, ejecutado con pytest",
            f"Total de Flujos: {len(FLUJOS)}",
            "Evidencia: video de la corrida, capturas por paso y logs del servidor",
        ],
    )
    est.tabla_indice(
        doc,
        "Índice de Flujos E2E",
        ["N°", "Flujo", "Archivo", "Estado"],
        [[f["num"], f["flujo"], f["archivo"], f["resultado"]] for f in FLUJOS],
    )

    for f in FLUJOS:
        tabla = est.nueva_tabla(doc)
        est.fila_titulo(tabla, f"Flujo {f['num']} — {f['flujo']}")
        est.fila_kv(tabla, "Archivo de prueba", f["archivo"])
        est.fila_kv(tabla, "Pantallas recorridas", f["pantallas"])
        est.fila_kv(tabla, "Precondición", f["precondicion"])
        est.fila_seccion(tabla, "PASOS")
        est.fila_kv(tabla, "Procedimiento", est.numerado(f["pasos"]))
        est.fila_seccion(tabla, "VERIFICACIÓN")
        est.fila_kv(tabla, "Comprobaciones", f["verificacion"])
        est.fila_kv(tabla, "Resultado", f["resultado"], fondo_valor=est.VERDE_OK)
        est.fila_seccion(tabla, "CONFIGURACIÓN")
        est.fila_kv(tabla, "Entorno", ENTORNO)
        est.fila_kv(tabla, "Comando", COMANDO)
        est.fila_kv(tabla, "Evidencia", EVIDENCIA_TXT)

        imagen = _captura(f["captura"])
        if imagen is not None:
            doc.add_paragraph()
            doc.add_picture(str(imagen), width=Inches(6.0))
            pie = doc.add_paragraph(f"Figura {f['num']}. {f['flujo']} (captura de la corrida).")
            pie.runs[0].italic = True

        doc.add_paragraph()
        if f["num"] != FLUJOS[-1]["num"]:
            doc.add_page_break()

    doc.save(str(destino))


def construir_md(destino: Path) -> None:
    L: list[str] = [
        "# Documento de Pruebas End to End (E2E)",
        "",
        f"**{est.SISTEMA}**",
        "",
        "- **Tipo de Prueba:** End to End sobre la aplicación completa (navegador real)",
        "- **Herramienta:** Selenium WebDriver 4 + Chrome, ejecutado con pytest",
        f"- **Total de Flujos:** {len(FLUJOS)}",
        f"- **{est.UNIVERSIDAD}**",
        "",
        "## Entorno",
        "",
        ENTORNO,
        "",
        f"Comando: `{COMANDO}`",
        "",
        "## Índice de Flujos E2E",
        "",
        "| N° | Flujo | Archivo | Estado |",
        "|----|-------|---------|--------|",
    ]
    for f in FLUJOS:
        L.append(f"| {f['num']} | {f['flujo']} | {f['archivo']} | **{f['resultado']}** |")
    L.append("")
    for f in FLUJOS:
        L.append(f"## Flujo {f['num']} — {f['flujo']}")
        L.append("")
        L.append(f"- **Archivo de prueba:** {f['archivo']}")
        L.append(f"- **Pantallas recorridas:** {f['pantallas']}")
        L.append(f"- **Precondición:** {f['precondicion']}")
        L.append("")
        L.append("**Pasos:**")
        L.append("")
        for i, paso in enumerate(f["pasos"], 1):
            L.append(f"{i}. {paso}")
        L.append("")
        L.append(f"- **Comprobaciones:** {f['verificacion']}")
        L.append(f"- **Resultado:** **{f['resultado']}**")
        imagen = _captura(f["captura"])
        if imagen is not None:
            ruta = imagen.relative_to(RAIZ).as_posix()
            L.append("")
            L.append(f"![Flujo {f['num']}](../{ruta})")
            L.append("")
            L.append(f"*Figura {f['num']}. {f['flujo']} (captura de la corrida).*")
        L.append("")
    L += [
        "## Evidencia de la corrida",
        "",
        "- `final/evidencias/pruebas/e2e_selenium/e2e_selenium.mp4` — video del recorrido "
        "completo en el navegador, con los logs del backend y del frontend al final.",
        "- `final/evidencias/pruebas/e2e_selenium/capturas/` — una captura por paso.",
        "- `final/evidencias/pruebas/e2e_selenium/logs/` — backend (uvicorn), frontend (Vite), "
        "consola del navegador y bitácora de pasos.",
        "- `final/evidencias/pruebas/e2e_selenium/pytest_e2e_selenium.txt` — salida de pytest.",
        "",
    ]
    destino.write_text("\n".join(L), encoding="utf-8")


def main() -> None:
    final = RAIZ / "final"
    construir_docx(final / "Pruebas End to End SPC.docx")
    construir_md(final / "Pruebas End to End SPC.md")
    print(f"OK  ->  {final / 'Pruebas End to End SPC.docx'}")
    print(f"OK  ->  {final / 'Pruebas End to End SPC.md'}")


if __name__ == "__main__":
    main()
