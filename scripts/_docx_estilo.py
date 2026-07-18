"""Estilo compartido de los documentos de prueba (.docx) del SPC.

Extrae la paleta y las utilidades de tabla del generador de caja negra
(``gen_pruebas_caja_negra.py``) para que los cuatro entregables —caja negra, pruebas
unitarias, experimentación y E2E— compartan un solo aspecto (misma portada, mismos
sombreados, misma tipografía) sin duplicar el estilo en cada script.

Requiere ``python-docx`` (``pip install python-docx``).
"""

from __future__ import annotations

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

# --- Paleta (idéntica al documento de caja negra de referencia) -----------------
AZUL_OSCURO = "1F4E79"   # cabecera de bloque y de la tabla índice
AZUL_MEDIO = "2E75B6"    # cabeceras de sección
AZUL_CLARO = "D6E4F0"    # celdas de etiqueta (columna izquierda)
VERDE_OK = "E2EFDA"      # celdas de resultado PASA
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)
GRIS_TXT = RGBColor(0x22, 0x22, 0x22)

UNIVERSIDAD = "Universidad Privada Antenor Orrego — Proyecto de Taller Integrador I"
SISTEMA = "SPC — Sistema de Predicción y Comercialización"


# --- Utilidades de estilo del .docx ---------------------------------------------
def sombrear(celda, hex_fill: str) -> None:
    """Rellena el fondo de una celda con un color hex."""
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    celda._tc.get_or_add_tcPr().append(shd)


def texto(celda, contenido: str, *, negrita: bool = False, blanco: bool = False, tam: int = 10) -> None:
    """Escribe texto en una celda respetando saltos de línea (``\\n``)."""
    celda.text = ""
    p = celda.paragraphs[0]
    lineas = str(contenido).split("\n")
    for i, linea in enumerate(lineas):
        run = p.add_run(linea)
        run.font.size = Pt(tam)
        run.font.bold = negrita
        run.font.color.rgb = BLANCO if blanco else GRIS_TXT
        if i < len(lineas) - 1:
            run.add_break()


def fila_titulo(tabla, titulo: str) -> None:
    """Fila-cabecera del bloque (azul oscuro, ocupa las dos columnas)."""
    fila = tabla.add_row()
    celda = fila.cells[0].merge(fila.cells[1])
    texto(celda, titulo, negrita=True, blanco=True, tam=11)
    sombrear(celda, AZUL_OSCURO)


def fila_seccion(tabla, titulo: str) -> None:
    """Fila de sección (azul medio, ocupa las dos columnas)."""
    fila = tabla.add_row()
    celda = fila.cells[0].merge(fila.cells[1])
    texto(celda, titulo, negrita=True, blanco=True)
    sombrear(celda, AZUL_MEDIO)


def fila_kv(tabla, etiqueta: str, valor: str, *, fondo_valor: str | None = None) -> None:
    """Fila etiqueta→valor (etiqueta en azul claro; valor opcionalmente sombreado)."""
    fila = tabla.add_row()
    texto(fila.cells[0], etiqueta, negrita=True)
    sombrear(fila.cells[0], AZUL_CLARO)
    texto(fila.cells[1], valor)
    if fondo_valor:
        sombrear(fila.cells[1], fondo_valor)


def vineta(items: list[str]) -> str:
    return "\n".join(f"• {x}" for x in items)


def numerado(items: list[str]) -> str:
    return "\n".join(f"{i}. {x}" for i, x in enumerate(items, 1))


# --- Bloques de documento reutilizables -----------------------------------------
def nuevo_documento() -> Document:
    """Documento con la tipografía base del proyecto."""
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)
    return doc


def portada(doc: Document, titulo: str, lineas: list[str]) -> None:
    """Portada centrada: título grande + subtítulo del sistema + líneas de metadatos."""
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run(titulo)
    r.bold = True
    r.font.size = Pt(20)
    r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rs = sub.add_run(SISTEMA)
    rs.bold = True
    rs.font.size = Pt(14)

    for linea in [*lineas, UNIVERSIDAD]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(linea).font.size = Pt(11)
    doc.add_paragraph()


def tabla_indice(doc: Document, titulo: str, cabeceras: list[str], filas: list[list[str]]) -> None:
    """Tabla índice con cabecera azul oscuro y una fila por elemento."""
    h = doc.add_paragraph()
    hr = h.add_run(titulo)
    hr.bold = True
    hr.font.size = Pt(13)

    tabla = doc.add_table(rows=1, cols=len(cabeceras))
    tabla.style = "Table Grid"
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    for celda, txt in zip(tabla.rows[0].cells, cabeceras):
        texto(celda, txt, negrita=True, blanco=True)
        sombrear(celda, AZUL_OSCURO)
    for fila in filas:
        celdas = tabla.add_row().cells
        for celda, valor in zip(celdas, fila):
            texto(celda, str(valor))
    doc.add_page_break()


def nueva_tabla(doc: Document):
    """Tabla de 2 columnas (etiqueta/valor) con rejilla, centrada."""
    tabla = doc.add_table(rows=0, cols=2)
    tabla.style = "Table Grid"
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    return tabla
